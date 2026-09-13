from openpyxl import load_workbook
wb=load_workbook('server/templates/formato-herramienta-danada.xlsx', data_only=False)
ws=wb.active
for row in ws.iter_rows():
    vals=[f'{c.coordinate}={c.value!r}' for c in row if c.value is not None]
    if vals: print(' | '.join(vals))
print('merged:', list(ws.merged_cells.ranges))
